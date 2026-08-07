#!/usr/bin/env python3
"""xlsx_extract.py — aisa process-capture: Layer 1 (deterministic extraction) + Layer 3 (replay).

Kernel asset (library/kernel/tools/). Read + executed at runtime, never edited at runtime.
Dependencies: Python 3.10+, openpyxl. No LLM, no network, no writes outside the given output paths.

Usage:
  L1 extract : python xlsx_extract.py <input.xlsx> <out.extraction.json> [--force] [--log <capture-log.md>]
  L3 replay  : python xlsx_extract.py --replay <input.xlsx> <extraction.json> <out.replay.md> [--log <capture-log.md>]

Contract (PROCESS_CAPTURE_SPEC.md):
- L1 extracts structure/logic only from what the file proves. Unreadable file -> status:"failed" JSON, never guessed.
- L1 caches on SHA-256: same hash as stored in an existing extraction JSON -> "cache-hit", no rewrite (unless --force).
- L3 is a fixed battery of checks, NOT a formula engine. No check = no claim; unsupported formulas -> "not replayable".
- Exit codes: 0 ok/cache-hit/failed-artefact-written, 2 usage error, 3 stale extraction on replay.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict

try:
    import openpyxl
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, column_index_from_string
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError:  # pragma: no cover
    print("xlsx_extract.py requires openpyxl (pip install openpyxl)", file=sys.stderr)
    sys.exit(2)

TOOL_VERSION = "1.0.0"
ARTEFACT_ID = "aisa.capture.extraction"

MAX_TOP_VALUES = 5
MAX_PATTERNS_PER_COLUMN = 5
MAX_EXCEPTION_CELLS = 20
MAX_DUP_VALUES = 20
MAX_CELLS_PER_DUP = 10
MAX_WHITESPACE_CELLS = 30
MAX_COMMENTS = 100
MAX_MISS_EXAMPLES = 5
MAX_OLDEST_ROWS = 5
MAX_SCAN_ROWS = 20_000          # per-sheet row cap for cell-level scans on huge sheets
BIG_SHEET_CELLS = 400_000
VALUE_TRUNC = 80

FIXED_AGE_BANDS = (30, 60, 90, 120, 180)
STALE_FILE_DAYS = 30

KEY_HEADER_RE = re.compile(r"(?i)(?:^|[^a-z])(id|key|code|ref|ticket|chave|c[oó]digo|n[ºo°]|num(?:ber|ero)?)(?:$|[^a-z])")
ARCHIVE_NAME_RE = re.compile(r"(?i)(closed|conclus|arquiv|hist[oó]r|archive|done|fechad)")
VOLATILE_RE = re.compile(r"(?i)\b(TODAY|NOW|RAND|RANDBETWEEN|RANDARRAY)\s*\(")
NOT_REPLAYABLE_RE = re.compile(r"(?i)\b(INDIRECT|OFFSET)\s*\(|\[\d+\]")  # dynamic refs / external workbook markers

_SHEETNAME = r"(?:'(?:[^']|'')+'|[A-Za-z_\u00C0-\u024F][\w.\u00C0-\u024F]{0,30})"
REF_RE = re.compile(
    rf"(?<![\w$:!.])(?:(?P<sheet>{_SHEETNAME})!)?"
    r"(?P<body>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}|\$?\d+:\$?\d+)"
    r"(?![\w(])"
)
STRING_SEG_RE = re.compile(r'"(?:[^"]|"")*"')
MAX_COL = 16384
MAX_ROW = 1048576


# ---------------------------------------------------------------- utilities

def utf8_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def now_iso() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


def json_default(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def short(value, limit: int = VALUE_TRUNC) -> str:
    text = value.isoformat() if isinstance(value, (dt.datetime, dt.date, dt.time)) else str(value)
    return text if len(text) <= limit else text[: limit - 1] + "…"


def append_log(log_path: str | None, layer: str, filename: str, event: str, detail: str) -> None:
    if not log_path:
        return
    new = not os.path.exists(log_path)
    with open(log_path, "a", encoding="utf-8") as fh:
        if new:
            fh.write("# Capture log\n\n| timestamp | layer | file | event | detail |\n|---|---|---|---|---|\n")
        fh.write(f"| {now_iso()} | {layer} | {filename} | {event} | {detail} |\n")


def cell_kind(value) -> str:
    if value is None or value == "":
        return "empty"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return "datetime"
    if isinstance(value, str):
        return "str"
    return "other"


def formula_text(cell) -> str | None:
    """Raw formula string of a cell in a formulas-workbook, or None."""
    if cell.data_type != "f":
        return None
    v = cell.value
    if isinstance(v, ArrayFormula):
        return v.text or ""
    return v if isinstance(v, str) else None


# ------------------------------------------------- A1 -> R1C1 normalization

def _split_strings(formula: str):
    """Yield (is_string, segment) pairs so refs inside string literals are never touched."""
    pos = 0
    for m in STRING_SEG_RE.finditer(formula):
        if m.start() > pos:
            yield False, formula[pos:m.start()]
        yield True, m.group(0)
        pos = m.end()
    if pos < len(formula):
        yield False, formula[pos:]


def _endpoint_r1c1(col: str | None, col_abs: bool, row: str | None, row_abs: bool,
                   base_row: int, base_col: int) -> str:
    parts = []
    if row is not None:
        r = int(row)
        if row_abs:
            parts.append(f"R{r}")
        else:
            d = r - base_row
            parts.append(f"R[{d}]" if d else "R")
    if col is not None:
        c = column_index_from_string(col)
        if col_abs:
            parts.append(f"C{c}")
        else:
            d = c - base_col
            parts.append(f"C[{d}]" if d else "C")
    return "".join(parts)


_ENDPOINT_RE = re.compile(r"(\$?)([A-Z]{1,3})?(\$?)(\d+)?")


def _ref_body_to_r1c1(body: str, base_row: int, base_col: int) -> str | None:
    """Convert one A1 ref body ('A2', '$A$1:$F$9', '$A:$F', '2:4') to relative R1C1. None = leave as-is."""
    out = []
    for endpoint in body.split(":"):
        m = _ENDPOINT_RE.fullmatch(endpoint)
        if not m:
            return None
        d1, col, d2, row = m.groups()
        if col is None and row is None:
            return None
        if col is not None and column_index_from_string(col) > MAX_COL:
            return None
        if row is not None and int(row) > MAX_ROW:
            return None
        # '$A' -> d1 before col; for row-only '$2' the '$' lands in d1 as well
        col_abs = bool(d1) if col is not None else False
        row_abs = bool(d2) if col is not None else bool(d1)
        out.append(_endpoint_r1c1(col, col_abs, row, row_abs, base_row, base_col))
    return ":".join(out)


def normalize_formula(formula: str, base_row: int, base_col: int) -> str:
    """Normalize every A1 reference to R1C1 relative form -> fill-down copies collapse to one pattern."""
    result = []
    for is_string, seg in _split_strings(formula):
        if is_string:
            result.append(seg)
            continue

        def repl(m: re.Match) -> str:
            converted = _ref_body_to_r1c1(m.group("body"), base_row, base_col)
            if converted is None:
                return m.group(0)
            sheet = m.group("sheet")
            return f"{sheet}!{converted}" if sheet else converted

        result.append(REF_RE.sub(repl, seg))
    return "".join(result)


def iter_refs(formula: str):
    """Yield (sheet_or_None, body) for every A1 reference outside string literals."""
    for is_string, seg in _split_strings(formula):
        if is_string:
            continue
        for m in REF_RE.finditer(seg):
            body = m.group("body")
            if _ref_body_to_r1c1(body, 1, 1) is None:
                continue
            sheet = m.group("sheet")
            if sheet and sheet.startswith("'"):
                sheet = sheet[1:-1].replace("''", "'")
            yield sheet, body


# ------------------------------------------------------------ L1 extraction

def detect_header(ws) -> tuple[int, int, str]:
    """Return (header_row, data_start_row, basis). Deterministic heuristic, basis recorded."""
    fp = ws.freeze_panes
    if fp:
        m = re.match(r"[A-Z]+(\d+)$", str(fp))
        if m and int(m.group(1)) >= 2:
            frozen_row = int(m.group(1))
            return frozen_row - 1, frozen_row, "frozen-pane"
    for r in range(1, min(ws.max_row or 1, 10) + 1):
        values = [c.value for c in ws[r]]
        nonempty = [v for v in values if v not in (None, "")]
        if len(nonempty) >= 2 and sum(isinstance(v, str) for v in nonempty) / len(nonempty) >= 0.6:
            return r, r + 1, "string-density"
    return 1, 2, "fallback-row-1"


def scan_columns(ws_data, ws_formula, header_row: int, data_start: int, data_end: int):
    """Per-column stats + formula patterns. Single pass over data rows on both workbook views."""
    max_col = ws_data.max_column or 0
    columns = []
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        header_val = ws_data.cell(row=header_row, column=col_idx).value
        header = str(header_val).strip() if header_val not in (None, "") else None

        kinds = Counter()
        values = Counter()
        nums, dates = [], []
        nulls = 0
        formulas: list[tuple[str, str]] = []          # (coord, raw formula)
        typed_cells: list[tuple[str, object]] = []    # (coord, value)

        for r in range(data_start, data_end + 1):
            dcell = ws_data.cell(row=r, column=col_idx)
            v = dcell.value
            kind = cell_kind(v)
            if kind == "empty":
                nulls += 1
            else:
                kinds[kind] += 1
                values[short(v)] += 1
                if kind == "number":
                    nums.append(v)
                elif kind == "datetime" and isinstance(v, (dt.datetime, dt.date)):
                    dates.append(v)
            raw = formula_text(ws_formula.cell(row=r, column=col_idx))
            if raw is not None:
                formulas.append((f"{letter}{r}", raw))
            elif kind != "empty":
                typed_cells.append((f"{letter}{r}", v))

        nonempty = sum(kinds.values())
        if header is None and nonempty == 0 and not formulas:
            continue

        if nonempty:
            ranked = kinds.most_common()
            inferred = ranked[0][0]
            mixed = len(ranked) > 1 and ranked[1][1] >= max(1, round(0.10 * nonempty))
            if mixed:
                inferred = f"mixed({ranked[0][0]},{ranked[1][0]})"
        else:
            inferred, mixed = "empty", False

        record = {
            "column": letter,
            "header": header,
            "inferred_type": inferred,
            "rows_nonempty": nonempty,
            "nulls": nulls,
            "distinct": len(values),
            "top_values": [{"value": v, "count": n} for v, n in values.most_common(MAX_TOP_VALUES)],
            "formula_count": len(formulas),
            "typed_count": len(typed_cells),
        }
        if nums:
            record["min"], record["max"] = min(nums), max(nums)
        if dates:
            record["min_date"], record["max_date"] = short(min(dates)), short(max(dates))
        record["_formulas"] = formulas
        record["_typed"] = typed_cells
        record["_mixed"] = mixed
        columns.append(record)
    return columns


def analyse_formulas(columns, data_start: int) -> None:
    """Group per-column formulas by R1C1 pattern; exceptions = manual overrides / divergent logic."""
    for col in columns:
        formulas = col.pop("_formulas")
        typed = col.pop("_typed")
        if not formulas:
            if typed:
                col["formula"] = None
            continue

        patterns = Counter()
        norm_by_cell = []
        for coord, raw in formulas:
            row = int(re.sub(r"[A-Z]+", "", coord))
            base_col = column_index_from_string(re.sub(r"\d+", "", coord))
            norm = normalize_formula(raw, row, base_col)
            patterns[norm] += 1
            norm_by_cell.append((coord, raw, norm))

        dominant, dom_count = patterns.most_common(1)[0]
        derived = len(formulas) >= max(1, len(typed))
        divergent = [
            {"cell": coord, "formula": short(raw, 160)}
            for coord, raw, norm in norm_by_cell if norm != dominant
        ][:MAX_EXCEPTION_CELLS]
        overrides = (
            [{"cell": coord, "value": short(v)} for coord, v in typed][:MAX_EXCEPTION_CELLS]
            if derived and typed else []
        )
        stray = (
            [{"cell": coord, "formula": short(raw, 160)} for coord, raw in formulas][:MAX_EXCEPTION_CELLS]
            if not derived else []
        )
        col["formula"] = {
            "count": len(formulas),
            "dominant_pattern": dominant,
            "dominant_count": dom_count,
            "coverage_of_formulas": round(dom_count / len(formulas), 4),
            "coverage_of_nonempty": round(dom_count / max(1, len(formulas) + len(typed)), 4),
            "patterns": [
                {"pattern": p, "count": n}
                for p, n in patterns.most_common(MAX_PATTERNS_PER_COLUMN)
            ],
            "other_pattern_count": max(0, len(patterns) - MAX_PATTERNS_PER_COLUMN),
            "exceptions": {
                "divergent_formulas": divergent,
                "typed_overrides": overrides,
                "formula_in_manual_column": stray,
            },
        }
        col["reads"] = sorted({
            (f"{sheet}!{body}" if sheet else body)
            for _, raw, _ in norm_by_cell
            for sheet, body in iter_refs(raw)
        })


def classify_columns(columns) -> None:
    """input / derived / manual per spec §4. Basis recorded so L2 can reason about it."""
    sheet_has_derived = any(
        c["formula_count"] >= max(1, c["typed_count"]) and c["formula_count"] > 0 for c in columns
    )
    for col in columns:
        if col["formula_count"] == 0 and col["rows_nonempty"] == 0:
            col["class"], col["class_basis"] = "empty", "no values, no formulas"
        elif col["formula_count"] >= max(1, col["typed_count"]):
            col["class"] = "derived"
            col["class_basis"] = f"{col['formula_count']} formula cells vs {col['typed_count']} typed"
        elif sheet_has_derived:
            col["class"] = "manual"
            col["class_basis"] = "typed values on a sheet with derived columns = human-maintained field"
        else:
            col["class"] = "input"
            col["class_basis"] = "typed values on a formula-free sheet = raw input/extract"


def is_key_like(col) -> bool:
    if col["class"] in ("derived", "empty"):
        return False
    if col["header"] and KEY_HEADER_RE.search(col["header"]):
        return True
    base = col["inferred_type"].startswith("str") or col["inferred_type"].startswith("mixed(str")
    return base and col["rows_nonempty"] >= 10 and col["distinct"] / col["rows_nonempty"] >= 0.95


def scan_anomalies(ws_data, columns, data_start: int, data_end: int):
    """Whitespace on key-like columns, mixed types, duplicate candidate keys."""
    anomalies = {"whitespace": [], "mixed_types": [], "duplicate_keys": []}
    for col in columns:
        col["key_like"] = is_key_like(col)
        if col["_mixed"]:
            anomalies["mixed_types"].append({
                "column": col["column"], "header": col["header"], "type": col["inferred_type"],
            })
        col.pop("_mixed", None)
        if not col["key_like"]:
            continue
        col_idx = column_index_from_string(col["column"])
        seen: dict[str, list[str]] = defaultdict(list)
        ws_cells = []
        for r in range(data_start, data_end + 1):
            v = ws_data.cell(row=r, column=col_idx).value
            if v in (None, ""):
                continue
            coord = f"{col['column']}{r}"
            if isinstance(v, str) and v != v.strip():
                ws_cells.append({"cell": coord, "value": repr(v)})
            seen[short(v, 200)].append(coord)
        if ws_cells:
            anomalies["whitespace"].append({
                "column": col["column"], "header": col["header"],
                "cells": ws_cells[:MAX_WHITESPACE_CELLS],
                "count": len(ws_cells),
            })
        dups = {v: cells for v, cells in seen.items() if len(cells) > 1}
        if dups:
            anomalies["duplicate_keys"].append({
                "column": col["column"], "header": col["header"],
                "values": [
                    {"value": v, "count": len(cells), "cells": cells[:MAX_CELLS_PER_DUP]}
                    for v, cells in sorted(dups.items())[:MAX_DUP_VALUES]
                ],
                "distinct_duplicated": len(dups),
            })
    return anomalies


def extract_validations(ws):
    out = []
    try:
        for dv in ws.data_validations.dataValidation:
            out.append({
                "range": str(dv.sqref),
                "type": dv.type,
                "operator": dv.operator,
                "formula1": dv.formula1,
                "formula2": dv.formula2,
                "allow_blank": dv.allowBlank,
            })
    except Exception as exc:  # never abort extraction on one group
        out.append({"error": f"validation read failed: {exc}"})
    return out


def extract_conditional_formatting(ws):
    out = []
    try:
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                fill_rgb = None
                try:
                    if rule.dxf is not None and rule.dxf.fill is not None:
                        color = rule.dxf.fill.bgColor or rule.dxf.fill.fgColor
                        fill_rgb = getattr(color, "rgb", None)
                        fill_rgb = fill_rgb if isinstance(fill_rgb, str) else None
                except Exception:
                    pass
                out.append({
                    "range": str(cf.sqref),
                    "type": rule.type,
                    "operator": rule.operator,
                    "formulas": list(rule.formula or []),
                    "priority": rule.priority,
                    "fill_rgb": fill_rgb,
                })
    except Exception as exc:
        out.append({"error": f"conditional formatting read failed: {exc}"})
    return out


def extract_fills(ws_data, columns, data_start: int, data_end: int, cf_rules):
    """Static cell fills per column (CF output is dynamic and never stored in cell.fill,
    so every static fill is a manual paint = candidate state encoding)."""
    cf_fill_rgbs = {r.get("fill_rgb") for r in cf_rules if isinstance(r, dict) and r.get("fill_rgb")}
    col_letters = {c["column"] for c in columns}
    fills: dict[str, Counter] = defaultdict(Counter)
    truncated = data_end - data_start + 1 > MAX_SCAN_ROWS
    end = min(data_end, data_start + MAX_SCAN_ROWS - 1)
    for row in ws_data.iter_rows(min_row=data_start, max_row=end):
        for cell in row:
            letter = cell.column_letter
            if letter not in col_letters:
                continue
            fill = cell.fill
            if fill is None or fill.patternType is None:
                continue
            rgb = getattr(fill.fgColor, "rgb", None)
            key = rgb if isinstance(rgb, str) else f"theme:{getattr(fill.fgColor, 'theme', '?')}"
            fills[letter][key] += 1
    out = []
    rows_by_letter = {c["column"]: max(1, c["rows_nonempty"]) for c in columns}
    for letter in sorted(fills):
        for rgb, count in fills[letter].most_common():
            out.append({
                "column": letter,
                "rgb": rgb,
                "count": count,
                "uniform": count >= 0.9 * rows_by_letter.get(letter, count),
                "matches_cf_output": rgb in cf_fill_rgbs,
            })
    return out, truncated


def extract_comments(ws):
    out = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.comment is not None:
                out.append({
                    "cell": cell.coordinate,
                    "author": cell.comment.author,
                    "text": short(cell.comment.text or "", 500),
                })
                if len(out) >= MAX_COMMENTS:
                    return out
    return out


def extract_named_ranges(wb):
    out = []
    try:
        for name, defn in wb.defined_names.items():
            out.append({"name": name, "target": defn.attr_text, "scope": "workbook"})
    except Exception as exc:
        out.append({"error": f"defined names read failed: {exc}"})
    for ws in wb.worksheets:
        try:
            for name, defn in getattr(ws, "defined_names", {}).items():
                out.append({"name": name, "target": defn.attr_text, "scope": ws.title})
        except Exception:
            continue
    return out


def extract_flags(path: str, wb, wb_data):
    flags = {"vba_present": False, "vba_modules": [], "external_links": [],
             "pivot_tables": [], "protected_sheets": []}
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            flags["vba_present"] = any(n.lower().endswith("vbaproject.bin") for n in names)
            for rel in (n for n in names if n.startswith("xl/externalLinks/_rels/")):
                for target in re.findall(rb'Target="([^"]+)"', zf.read(rel)):
                    flags["external_links"].append(target.decode("utf-8", "replace"))
    except Exception:
        pass
    for ws in wb_data.worksheets:
        try:
            if ws.protection and ws.protection.sheet:
                flags["protected_sheets"].append(ws.title)
        except Exception:
            continue
        try:
            for pivot in getattr(ws, "_pivots", []):
                src = pivot.cache.cacheSource.worksheetSource
                flags["pivot_tables"].append({
                    "sheet": ws.title,
                    "source": f"{getattr(src, 'sheet', '?')}!{getattr(src, 'ref', '?')}",
                })
        except Exception:
            continue
    return flags


def extract(path: str, out_path: str, force: bool, log_path: str | None) -> int:
    filename = os.path.basename(path)
    if not os.path.exists(path):
        print(f"input not found: {path}", file=sys.stderr)
        return 2

    sha = sha256_file(path)
    if not force and os.path.exists(out_path):
        try:
            with open(out_path, encoding="utf-8") as fh:
                previous = json.load(fh)
            if previous.get("identity", {}).get("sha256") == sha and previous.get("status") == "ok":
                print(f"cache-hit: {filename} unchanged (sha256 {sha[:8]}), extraction kept")
                append_log(log_path, "L1", filename, "cache-hit", f"sha256 {sha[:8]}")
                return 0
        except Exception:
            pass  # unreadable previous artefact -> re-extract

    stat = os.stat(path)
    identity = {
        "filename": filename,
        "source_path": path,
        "size_bytes": stat.st_size,
        "sha256": sha,
        "modified": dt.datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
    }
    doc = {
        "artefact": ARTEFACT_ID,
        "tool": {"name": "xlsx_extract.py", "version": TOOL_VERSION,
                 "openpyxl": openpyxl.__version__,
                 "python": ".".join(map(str, sys.version_info[:3]))},
        "extracted_at": now_iso(),
        "identity": identity,
    }

    try:
        wb_formula = load_workbook(path, data_only=False)
        wb_data = load_workbook(path, data_only=True)
    except Exception as exc:
        doc.update({"status": "failed", "reason": f"{type(exc).__name__}: {exc}"})
        _write_json(out_path, doc)
        print(f"extraction FAILED for {filename}: {exc} -> failure artefact written")
        append_log(log_path, "L1", filename, "failed", short(str(exc), 120))
        return 0

    doc["status"] = "ok"
    doc["workbook"] = {
        "named_ranges": extract_named_ranges(wb_formula),
        "flags": extract_flags(path, wb_formula, wb_data),
    }

    sheets = []
    known_sheets = set(wb_data.sheetnames)
    for ws_data, ws_formula in zip(wb_data.worksheets, wb_formula.worksheets):
        errors = []
        header_row, data_start, basis = detect_header(ws_data)
        data_end = ws_data.max_row or header_row
        sheet = {
            "name": ws_data.title,
            "state": ws_data.sheet_state,
            "dimensions": ws_data.dimensions,
            "max_row": ws_data.max_row,
            "max_col": ws_data.max_column,
            "frozen_panes": str(ws_data.freeze_panes) if ws_data.freeze_panes else None,
            "header_row": header_row,
            "data_start_row": data_start,
            "data_end_row": data_end,
            "header_basis": basis,
        }
        try:
            columns = scan_columns(ws_data, ws_formula, header_row, data_start, data_end)
            analyse_formulas(columns, data_start)
            classify_columns(columns)
            sheet["anomalies"] = scan_anomalies(ws_data, columns, data_start, data_end)
        except Exception as exc:
            columns = []
            sheet["anomalies"] = {}
            errors.append(f"column scan failed: {type(exc).__name__}: {exc}")
        sheet["validations"] = extract_validations(ws_formula)
        sheet["conditional_formatting"] = extract_conditional_formatting(ws_formula)
        try:
            sheet["fills"], truncated = extract_fills(
                ws_data, columns, data_start, data_end, sheet["conditional_formatting"])
            if truncated:
                sheet["fills_truncated_at_rows"] = MAX_SCAN_ROWS
        except Exception as exc:
            sheet["fills"] = []
            errors.append(f"fill scan failed: {type(exc).__name__}: {exc}")
        sheet["comments"] = extract_comments(ws_formula)
        sheet["columns"] = columns
        unknown = sorted({
            ref.split("!")[0]
            for col in columns for ref in col.get("reads", [])
            if "!" in ref and ref.split("!")[0] not in known_sheets
        })
        if unknown:
            sheet["unknown_sheet_refs"] = unknown
        if errors:
            sheet["errors"] = errors
        sheets.append(sheet)
    doc["sheets"] = sheets

    _write_json(out_path, doc)
    n_cols = sum(len(s["columns"]) for s in sheets)
    print(f"extracted: {filename} -> {out_path} ({len(sheets)} sheets, {n_cols} columns, sha256 {sha[:8]})")
    append_log(log_path, "L1", filename, "extracted", f"{len(sheets)} sheets, {n_cols} columns, sha256 {sha[:8]}")
    return 0


def _write_json(out_path: str, doc: dict) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    tmp = out_path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, ensure_ascii=False, indent=1, default=json_default)
    os.replace(tmp, out_path)


# --------------------------------------------------------------- L3 replay

class Finding:
    __slots__ = ("check", "location", "expected", "found", "severity")
    ORDER = {"high": 0, "medium": 1, "low": 2, "info": 3}

    def __init__(self, check, location, expected, found, severity):
        self.check, self.location = check, location
        self.expected, self.found, self.severity = expected, found, severity


def find_calls(text: str, fname: str) -> list[str]:
    """Balanced-paren extraction of every `FNAME(...)` call, string-literal aware."""
    out = []
    upper = text.upper()
    i = 0
    while True:
        j = upper.find(fname + "(", i)
        if j == -1:
            return out
        if j > 0 and (text[j - 1].isalnum() or text[j - 1] in "_.$"):
            i = j + 1
            continue
        depth, in_str = 0, False
        for k in range(j + len(fname), len(text)):
            ch = text[k]
            if in_str:
                if ch == '"':
                    in_str = False
            elif ch == '"':
                in_str = True
            elif ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0:
                    out.append(text[j:k + 1])
                    break
        i = j + 1


def split_args(call: str) -> list[str]:
    inner = call[call.index("(") + 1:-1]
    args, depth, in_str, start = [], 0, False, 0
    for k, ch in enumerate(inner):
        if in_str:
            if ch == '"':
                in_str = False
        elif ch == '"':
            in_str = True
        elif ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        elif ch == "," and depth == 0:
            args.append(inner[start:k].strip())
            start = k + 1
    args.append(inner[start:].strip())
    return args


SINGLE_CELL_RE = re.compile(rf"^(?:(?P<sheet>{_SHEETNAME})!)?\$?(?P<col>[A-Z]{{1,3}})\$?(?P<row>\d+)$")
RANGE_ARG_RE = re.compile(
    rf"^(?:(?P<sheet>{_SHEETNAME})!)?"
    r"(?:\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)?:\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+)?)$"
)


def _unquote_sheet(sheet: str | None) -> str | None:
    if sheet and sheet.startswith("'"):
        return sheet[1:-1].replace("''", "'")
    return sheet


class Replayer:
    """Fixed check battery over one workbook + its extraction JSON. No general evaluation."""

    def __init__(self, wb_data, extraction: dict):
        self.wb = wb_data
        self.x = extraction
        self.range_cache: dict[tuple, list[tuple[int, object]]] = {}
        self.findings: list[Finding] = []
        self.not_replayable: Counter = Counter()
        self.nr_example: dict[str, str] = {}
        self.lookup_targets: set[tuple[str, str]] = set()   # (sheet, col letter) used as match column
        self.criterion_cols: set[tuple[str, str]] = set()   # (sheet, col letter) feeding lookup values
        self.checked_cells = 0
        self.skipped_empty = 0

    # -- resolution helpers ------------------------------------------------
    def sheet_meta(self, name: str) -> dict | None:
        for s in self.x["sheets"]:
            if s["name"] == name:
                return s
        return None

    def resolve_cell(self, default_sheet: str, arg: str):
        m = SINGLE_CELL_RE.fullmatch(arg)
        if not m:
            return None, None
        sheet = _unquote_sheet(m.group("sheet")) or default_sheet
        if sheet not in self.wb.sheetnames:
            return None, None
        coord = (sheet, m.group("col"), int(m.group("row")))
        ws = self.wb[sheet]
        return coord, ws.cell(row=coord[2], column=column_index_from_string(coord[1])).value

    def materialize(self, default_sheet: str, arg: str):
        """Range arg -> (sheet, first_col_idx, width, [(row, first_col_value)...]) or None."""
        m = RANGE_ARG_RE.fullmatch(arg)
        if not m:
            sc = SINGLE_CELL_RE.fullmatch(arg)
            if sc:  # single-cell "range"
                sheet = _unquote_sheet(sc.group("sheet")) or default_sheet
                if sheet not in self.wb.sheetnames:
                    return None
                c = column_index_from_string(sc.group("col"))
                r = int(sc.group("row"))
                return sheet, c, 1, [(r, self.wb[sheet].cell(row=r, column=c).value)]
            return None
        sheet = _unquote_sheet(m.group("sheet")) or default_sheet
        if sheet not in self.wb.sheetnames:
            return None
        ws = self.wb[sheet]
        c1 = column_index_from_string(m.group("c1"))
        c2 = column_index_from_string(m.group("c2"))
        if c2 < c1:
            c1, c2 = c2, c1
        r1 = int(m.group("r1")) if m.group("r1") else 1
        r2 = int(m.group("r2")) if m.group("r2") else (ws.max_row or 1)
        r2 = min(r2, ws.max_row or 1)
        key = (sheet, c1, r1, r2)
        if key not in self.range_cache:
            self.range_cache[key] = [
                (r, ws.cell(row=r, column=c1).value) for r in range(r1, r2 + 1)
            ]
        return sheet, c1, c2 - c1 + 1, self.range_cache[key]

    @staticmethod
    def _eq(a, b, trim: bool) -> bool:
        if isinstance(a, str) and isinstance(b, str):
            if trim:
                a, b = a.strip(), b.strip()
            return a.casefold() == b.casefold()   # Excel text matching is case-insensitive
        if isinstance(a, bool) or isinstance(b, bool):
            return a is b
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            return abs(a - b) < 1e-9
        return a == b

    def match_row(self, rows, needle, trim: bool):
        for r, v in rows:
            if v in (None, ""):
                continue
            if self._eq(v, needle, trim):
                return r
        return None

    @staticmethod
    def values_equal(a, b) -> bool:
        if a in (None, "") and b in (None, ""):
            return True
        if isinstance(a, (int, float)) and isinstance(b, (int, float)) \
                and not isinstance(a, bool) and not isinstance(b, bool):
            return abs(a - b) < 1e-6
        if isinstance(a, (dt.datetime, dt.date)) and isinstance(b, (dt.datetime, dt.date)):
            return str(a) == str(b)
        return a == b

    # -- battery checks ----------------------------------------------------
    def run(self) -> None:
        self.check_lookups()
        self.check_duplicates()
        self.check_whitespace_casing()
        self.check_staleness()
        self.check_pattern_exceptions()
        self.check_orphans()

    def _mark_not_replayable(self, sheet_name, coord, raw, reason) -> bool:
        pattern = f"{reason}: {short(raw, 100)}"
        self.not_replayable[pattern] += 1
        self.nr_example.setdefault(pattern, f"{sheet_name}!{coord}")
        return True

    def check_lookups(self) -> None:
        misses_by_col: dict[str, list] = defaultdict(list)
        stale_by_col: dict[str, list] = defaultdict(list)
        for sheet in self.x["sheets"]:
            name = sheet["name"]
            if name not in self.wb.sheetnames:
                continue
            ws = self.wb[name]
            for col in sheet["columns"]:
                if not col.get("formula"):
                    continue
                pattern = col["formula"]["dominant_pattern"]
                has_lookup = any(t in pattern.upper() for t in ("VLOOKUP(", "XLOOKUP(", "COUNTIF(", "INDEX("))
                if not has_lookup:
                    continue
                for r in range(sheet["data_start_row"], sheet["data_end_row"] + 1):
                    coord = f"{col['column']}{r}"
                    fcell = self._formula_at(name, coord)
                    if fcell is None:
                        continue
                    if NOT_REPLAYABLE_RE.search(fcell):
                        self._mark_not_replayable(name, coord, fcell, "dynamic/external reference")
                        continue
                    stored = ws[coord].value
                    volatile = bool(VOLATILE_RE.search(fcell))
                    self._replay_cell(name, coord, fcell, stored, volatile, misses_by_col, stale_by_col)
        for key, misses in sorted(misses_by_col.items()):
            examples = ", ".join(f"{short(v, 30)} @{c}" for c, v in misses[:MAX_MISS_EXAMPLES])
            self.findings.append(Finding(
                "lookup integrity", key,
                "every lookup value resolvable in its target range",
                f"{len(misses)} lookup value(s) with no match in target (e.g. {examples})",
                "medium"))
        for key, stales in sorted(stale_by_col.items()):
            examples = "; ".join(f"{c}: stored {short(s, 24)} vs recomputed {short(rc, 24)}"
                                 for c, s, rc in stales[:MAX_MISS_EXAMPLES])
            self.findings.append(Finding(
                "lookup integrity", key,
                "stored result equals recomputed lookup result",
                f"{len(stales)} cell(s) diverge ({examples})",
                "medium"))

    def _formula_at(self, sheet_name: str, coord: str) -> str | None:
        cache = getattr(self, "_fwb", None)
        if cache is None:
            return None
        try:
            return formula_text(cache[sheet_name][coord])
        except Exception:
            return None

    def _replay_cell(self, sheet_name, coord, formula, stored, volatile, misses_by_col, stale_by_col):
        col_letter = re.sub(r"[0-9]+$", "", coord)
        col_key = f"{sheet_name}!{col_letter} (VLOOKUP)"
        handled = False
        for call in find_calls(formula, "VLOOKUP"):
            args = split_args(call)
            if len(args) < 4 or args[3].upper().rstrip(")") not in ("0", "FALSE"):
                self._mark_not_replayable(sheet_name, coord, call, "approximate-match VLOOKUP")
                continue
            needle = self._resolve_scalar(sheet_name, args[0], coord)
            if needle is _UNSUPPORTED:
                self._mark_not_replayable(sheet_name, coord, call, "unsupported lookup value")
                continue
            table = self.materialize(sheet_name, args[1])
            if table is None:
                self._mark_not_replayable(sheet_name, coord, call, "unsupported table range")
                continue
            t_sheet, t_first, width, rows = table
            try:
                idx = int(args[2])
            except ValueError:
                self._mark_not_replayable(sheet_name, coord, call, "non-literal column index")
                continue
            self.lookup_targets.add((t_sheet, get_column_letter(t_first)))
            self._note_criterion(sheet_name, args[0])
            if idx > width:
                self.findings.append(Finding(
                    "orphan references", f"{sheet_name}!{coord}",
                    f"VLOOKUP column index within range width {width}",
                    f"index {idx} exceeds {args[1]}", "high"))
                continue
            handled = True
            if needle in (None, ""):
                self.skipped_empty += 1
                continue
            self.checked_cells += 1
            raw_row = self.match_row(rows, needle, trim=False)
            if raw_row is None:
                norm_row = self.match_row(rows, needle, trim=True) if isinstance(needle, str) else None
                if norm_row is not None:
                    tcell = f"{t_sheet}!{get_column_letter(t_first)}{norm_row}"
                    tval = self.wb[t_sheet].cell(row=norm_row, column=t_first).value
                    self.findings.append(Finding(
                        "whitespace/casing", f"{sheet_name}!{coord}",
                        f"lookup of {short(needle, 40)!r} matches {tcell}",
                        f"raw match fails; succeeds only after TRIM ({tcell} holds {short(tval, 40)!r}) "
                        "— silent lookup failure", "high"))
                else:
                    misses_by_col[col_key].append((coord, needle))
            elif not volatile:
                target_val = self.wb[t_sheet].cell(row=raw_row, column=t_first + idx - 1).value
                if not self.values_equal(stored, target_val):
                    stale_by_col[col_key].append((coord, stored, target_val))
        for call in find_calls(formula, "COUNTIF"):
            args = split_args(call)
            if len(args) != 2:
                self._mark_not_replayable(sheet_name, coord, call, "unsupported COUNTIF arity")
                continue
            needle = self._resolve_scalar(sheet_name, args[1], coord)
            if needle is _UNSUPPORTED or (isinstance(needle, str) and needle.startswith(("<", ">", "="))):
                self._mark_not_replayable(sheet_name, coord, call, "unsupported COUNTIF criterion")
                continue
            rng = self.materialize(sheet_name, args[0])
            if rng is None:
                self._mark_not_replayable(sheet_name, coord, call, "unsupported COUNTIF range")
                continue
            t_sheet, t_first, _, rows = rng
            self.lookup_targets.add((t_sheet, get_column_letter(t_first)))
            self._note_criterion(sheet_name, args[1])
            handled = True
            if needle in (None, ""):
                self.skipped_empty += 1
                continue
            self.checked_cells += 1
            raw_n = sum(1 for _, v in rows if v not in (None, "") and self._eq(v, needle, trim=False))
            if raw_n == 0 and isinstance(needle, str):
                norm_n = sum(1 for _, v in rows if v not in (None, "") and self._eq(v, needle, trim=True))
                if norm_n > 0:
                    self.findings.append(Finding(
                        "whitespace/casing", f"{sheet_name}!{coord}",
                        f"membership of {short(needle, 40)!r} in {args[0]}",
                        "raw COUNTIF = 0 but > 0 after TRIM — whitespace breaks the match", "high"))
                else:
                    misses_by_col[f"{sheet_name}!{col_letter} (COUNTIF membership)"].append((coord, needle))
        for fname in ("XLOOKUP", "INDEX"):
            for call in find_calls(formula, fname):
                if fname == "INDEX" and "MATCH(" not in call.upper():
                    continue
                handled_x = self._replay_xlookup(sheet_name, coord, call, stored, volatile, stale_by_col, misses_by_col) \
                    if fname == "XLOOKUP" else \
                    self._replay_index_match(sheet_name, coord, call, stored, volatile, stale_by_col, misses_by_col)
                handled = handled or handled_x
        if not handled and not find_calls(formula, "VLOOKUP") and not find_calls(formula, "COUNTIF"):
            pass  # column matched lookup keywords only via dominant pattern; nothing claimable per-cell

    def _note_criterion(self, default_sheet: str, arg: str) -> None:
        m = SINGLE_CELL_RE.fullmatch(arg)
        if m:
            sheet = _unquote_sheet(m.group("sheet")) or default_sheet
            self.criterion_cols.add((sheet, m.group("col")))

    def _resolve_scalar(self, default_sheet: str, arg: str, coord: str):
        arg = arg.strip()
        if arg.startswith('"') and arg.endswith('"'):
            return arg[1:-1].replace('""', '"')
        try:
            return float(arg) if "." in arg else int(arg)
        except ValueError:
            pass
        ref, val = self.resolve_cell(default_sheet, arg)
        if ref is not None:
            return val
        return _UNSUPPORTED

    def _replay_xlookup(self, sheet_name, coord, call, stored, volatile, stale_by_col, misses_by_col) -> bool:
        args = split_args(call)
        if len(args) < 3:
            return self._mark_not_replayable(sheet_name, coord, call, "unsupported XLOOKUP arity") and False
        if len(args) >= 5 and args[4] not in ("0", ""):
            return self._mark_not_replayable(sheet_name, coord, call, "non-exact XLOOKUP mode") and False
        needle = self._resolve_scalar(sheet_name, args[0], coord)
        lookup_rng = self.materialize(sheet_name, args[1])
        return_rng = self.materialize(sheet_name, args[2])
        if needle is _UNSUPPORTED or lookup_rng is None or return_rng is None:
            return self._mark_not_replayable(sheet_name, coord, call, "unsupported XLOOKUP args") and False
        l_sheet, l_first, _, rows = lookup_rng
        r_sheet, r_first, _, _ = return_rng
        self.lookup_targets.add((l_sheet, get_column_letter(l_first)))
        self._note_criterion(sheet_name, args[0])
        if needle in (None, ""):
            self.skipped_empty += 1
            return True
        self.checked_cells += 1
        raw_row = self.match_row(rows, needle, trim=False)
        col_key = f"{sheet_name}!{re.sub(r'[0-9]+$', '', coord)} (XLOOKUP)"
        if raw_row is None:
            norm_row = self.match_row(rows, needle, trim=True) if isinstance(needle, str) else None
            if norm_row is not None:
                self.findings.append(Finding(
                    "whitespace/casing", f"{sheet_name}!{coord}",
                    f"XLOOKUP of {short(needle, 40)!r} matches row {norm_row}",
                    "raw match fails; succeeds only after TRIM — silent lookup failure", "high"))
            else:
                misses_by_col[col_key].append((coord, needle))
        elif not volatile:
            got = self.wb[r_sheet].cell(row=raw_row, column=r_first).value
            if not self.values_equal(stored, got):
                stale_by_col[col_key].append((coord, stored, got))
        return True

    def _replay_index_match(self, sheet_name, coord, call, stored, volatile, stale_by_col, misses_by_col) -> bool:
        args = split_args(call)
        if len(args) < 2:
            return False
        matches = find_calls(args[1], "MATCH")
        if not matches:
            return self._mark_not_replayable(sheet_name, coord, call, "INDEX without literal MATCH") and False
        margs = split_args(matches[0])
        if len(margs) < 3 or margs[2] not in ("0",):
            return self._mark_not_replayable(sheet_name, coord, call, "non-exact MATCH") and False
        needle = self._resolve_scalar(sheet_name, margs[0], coord)
        match_rng = self.materialize(sheet_name, margs[1])
        index_rng = self.materialize(sheet_name, args[0])
        if needle is _UNSUPPORTED or match_rng is None or index_rng is None:
            return self._mark_not_replayable(sheet_name, coord, call, "unsupported INDEX/MATCH args") and False
        m_sheet, m_first, _, rows = match_rng
        i_sheet, i_first, _, _ = index_rng
        self.lookup_targets.add((m_sheet, get_column_letter(m_first)))
        self._note_criterion(sheet_name, margs[0])
        if needle in (None, ""):
            self.skipped_empty += 1
            return True
        self.checked_cells += 1
        raw_row = self.match_row(rows, needle, trim=False)
        col_key = f"{sheet_name}!{re.sub(r'[0-9]+$', '', coord)} (INDEX+MATCH)"
        if raw_row is None:
            norm_row = self.match_row(rows, needle, trim=True) if isinstance(needle, str) else None
            if norm_row is not None:
                self.findings.append(Finding(
                    "whitespace/casing", f"{sheet_name}!{coord}",
                    f"MATCH of {short(needle, 40)!r} finds row {norm_row}",
                    "raw match fails; succeeds only after TRIM — silent lookup failure", "high"))
            else:
                misses_by_col[col_key].append((coord, needle))
        elif not volatile:
            offset = raw_row - rows[0][0]
            got = self.wb[i_sheet].cell(row=index_rng[3][0][0] + offset, column=i_first).value
            if not self.values_equal(stored, got):
                stale_by_col[col_key].append((coord, stored, got))
        return True

    def check_duplicates(self) -> None:
        for sheet in self.x["sheets"]:
            for dup_col in sheet.get("anomalies", {}).get("duplicate_keys", []):
                col_letter = dup_col["column"]
                is_target = (sheet["name"], col_letter) in self.lookup_targets
                for entry in dup_col["values"]:
                    diff_note = self._duplicate_row_diff(sheet, col_letter, entry["cells"])
                    self.findings.append(Finding(
                        "key uniqueness",
                        f"{sheet['name']}!{', '.join(entry['cells'])}",
                        f"'{entry['value']}' unique in {sheet['name']}!{col_letter} "
                        f"({dup_col.get('header') or 'key column'}"
                        f"{'; lookup/COUNTIF target — first match wins' if is_target else ''})",
                        f"{entry['count']}x present{diff_note}",
                        "high" if is_target else "medium"))

    def _duplicate_row_diff(self, sheet, col_letter, cells) -> str:
        if sheet["name"] not in self.wb.sheetnames or len(cells) < 2:
            return ""
        ws = self.wb[sheet["name"]]
        rows = [int(re.sub(r"[A-Z]+", "", c)) for c in cells[:2]]
        differing = []
        for col in sheet["columns"]:
            if col["column"] == col_letter:
                continue
            idx = column_index_from_string(col["column"])
            a = ws.cell(row=rows[0], column=idx).value
            b = ws.cell(row=rows[1], column=idx).value
            if not self.values_equal(a, b):
                differing.append(col.get("header") or col["column"])
        if differing:
            return f"; rows differ in: {', '.join(differing[:6])} — conflicting records"
        return "; rows identical — redundant record"

    def check_whitespace_casing(self) -> None:
        for sheet in self.x["sheets"]:
            for ws_col in sheet.get("anomalies", {}).get("whitespace", []):
                key = (sheet["name"], ws_col["column"])
                in_lookups = key in self.lookup_targets or key in self.criterion_cols
                for cell in ws_col["cells"]:
                    self.findings.append(Finding(
                        "whitespace/casing",
                        f"{sheet['name']}!{cell['cell']}",
                        f"key without leading/trailing whitespace in "
                        f"'{ws_col.get('header') or ws_col['column']}'",
                        f"{cell['value']}"
                        + (" — column feeds lookups/COUNTIF: raw matching against this key is impossible"
                           if in_lookups else ""),
                        "high" if in_lookups else "low"))
            # casefold collisions on key-like columns
            if sheet["name"] not in self.wb.sheetnames:
                continue
            ws = self.wb[sheet["name"]]
            for col in sheet["columns"]:
                if not col.get("key_like"):
                    continue
                idx = column_index_from_string(col["column"])
                groups: dict[str, set] = defaultdict(set)
                for r in range(sheet["data_start_row"], sheet["data_end_row"] + 1):
                    v = ws.cell(row=r, column=idx).value
                    if isinstance(v, str) and v.strip():
                        groups[v.strip().casefold()].add(v.strip())
                for folded, variants in sorted(groups.items()):
                    if len(variants) > 1:
                        self.findings.append(Finding(
                            "whitespace/casing",
                            f"{sheet['name']}!{col['column']} "
                            f"({col.get('header') or 'key column'})",
                            "one canonical casing per key",
                            f"case variants coexist: {sorted(variants)}",
                            "low"))

    def _cf_thresholds(self) -> dict[tuple[str, str], set[int]]:
        out: dict[tuple[str, str], set[int]] = defaultdict(set)
        for sheet in self.x["sheets"]:
            for rule in sheet.get("conditional_formatting", []):
                if not isinstance(rule, dict) or rule.get("error"):
                    continue
                cols = set(re.findall(r"\$?([A-Z]{1,3})\$?\d+", str(rule.get("range", ""))))
                for f in rule.get("formulas", []):
                    for n in re.findall(r"[<>]=?\s*(\d+)", str(f)):
                        for c in cols:
                            out[(sheet["name"], c)].add(int(n))
        return out

    def check_staleness(self) -> None:
        today = dt.date.today()
        cf_thresholds = self._cf_thresholds()
        all_cf_days = sorted({d for days in cf_thresholds.values() for d in days if 0 < d <= 400})
        for sheet in self.x["sheets"]:
            if sheet["name"] not in self.wb.sheetnames:
                continue
            ws = self.wb[sheet["name"]]
            archive_like = bool(ARCHIVE_NAME_RE.search(sheet["name"]))
            key_col_idx = next(
                (column_index_from_string(c["column"]) for c in sheet["columns"] if c.get("key_like")),
                None)
            for col in sheet["columns"]:
                if "datetime" not in col["inferred_type"]:
                    continue
                header = col.get("header") or col["column"]
                closing_like = bool(re.search(r"(?i)(conclus|clos|fecho|fim|end)", header))
                idx = column_index_from_string(col["column"])
                ages = []
                for r in range(sheet["data_start_row"], sheet["data_end_row"] + 1):
                    v = ws.cell(row=r, column=idx).value
                    if isinstance(v, (dt.datetime, dt.date)):
                        d = v.date() if isinstance(v, dt.datetime) else v
                        key_val = (ws.cell(row=r, column=key_col_idx).value
                                   if key_col_idx else None)
                        ages.append(((today - d).days, f"{col['column']}{r}", key_val))
                if not ages:
                    continue
                bands = sorted(set(FIXED_AGE_BANDS) | set(all_cf_days))
                dist = {f">={b}d": sum(1 for a, _, _ in ages if a >= b) for b in bands}
                future = sum(1 for a, _, _ in ages if a < 0)
                oldest = sorted(ages, reverse=True)[:MAX_OLDEST_ROWS]
                oldest_txt = "; ".join(
                    f"{a}d @{c}" + (f" ({short(k, 24)})" if k not in (None, "") else "")
                    for a, c, k in oldest)
                max_cf = max(all_cf_days) if all_cf_days else None
                beyond_cf = dist.get(f">={max_cf}d") if max_cf else None
                severity = "info"
                if not archive_like and not closing_like:
                    if dist.get(">=120d"):
                        severity = "medium"
                    elif beyond_cf:
                        severity = "medium"
                found = (f"{len(ages)} dated rows; distribution {dist}"
                         + (f"; {future} in the future" if future else "")
                         + f"; oldest: {oldest_txt}")
                expected = ("rows within the alerting thresholds present in conditional formatting "
                            f"(max {max_cf}d)" if max_cf else "recent activity")
                self.findings.append(Finding(
                    "staleness",
                    f"{sheet['name']}!{col['column']} ({header})",
                    expected, found, severity))
        # file-level staleness
        try:
            mtime = dt.datetime.fromisoformat(self.x["identity"]["modified"]).date()
            age = (dt.date.today() - mtime).days
            if age > STALE_FILE_DAYS:
                self.findings.append(Finding(
                    "staleness", self.x["identity"]["filename"],
                    f"file touched within {STALE_FILE_DAYS} days",
                    f"last modified {mtime.isoformat()} ({age} days ago) — stored TODAY()-based "
                    "values are frozen at an older date", "info"))
        except Exception:
            pass

    @staticmethod
    def _cell_span(cells: list[str]) -> str:
        rows = sorted(int(re.sub(r"[A-Z]+", "", c)) for c in cells)
        letter = re.sub(r"[0-9]+", "", cells[0])
        if len(rows) == 1:
            return f"{letter}{rows[0]}"
        return f"{letter}{rows[0]}:{letter}{rows[-1]} ({len(cells)} cells)"

    def check_pattern_exceptions(self) -> None:
        for sheet in self.x["sheets"]:
            for col in sheet["columns"]:
                f = col.get("formula")
                if not f:
                    continue
                exc = f.get("exceptions", {})
                loc = f"{sheet['name']}!{col['column']}"
                overrides = exc.get("typed_overrides", [])
                if overrides:
                    examples = "; ".join(f"{e['cell']}={e['value']!r}" for e in overrides[:3])
                    self.findings.append(Finding(
                        "pattern exceptions",
                        f"{sheet['name']}!{self._cell_span([e['cell'] for e in overrides])}",
                        f"formula column ({short(f['dominant_pattern'], 60)})",
                        f"{len(overrides)} typed value(s) replace the formula — manual override "
                        f"of automated logic ({examples})", "medium"))
                divergent = exc.get("divergent_formulas", [])
                if divergent:
                    families = len({e["formula"] for e in divergent})
                    examples = "; ".join(f"{e['cell']}: {short(e['formula'], 60)}" for e in divergent[:2])
                    self.findings.append(Finding(
                        "pattern exceptions",
                        f"{sheet['name']}!{self._cell_span([e['cell'] for e in divergent])}",
                        f"one fill-down pattern for the whole column "
                        f"({short(f['dominant_pattern'], 60)})",
                        f"{len(divergent)} cell(s) carry {families} divergent formula variant(s) — "
                        f"manually re-anchored copies can misalign rows ({examples})", "medium"))
                stray = exc.get("formula_in_manual_column", [])
                if stray:
                    examples = "; ".join(f"{e['cell']}: {short(e['formula'], 60)}" for e in stray[:2])
                    self.findings.append(Finding(
                        "pattern exceptions",
                        f"{sheet['name']}!{self._cell_span([e['cell'] for e in stray])}",
                        f"typed column {loc}",
                        f"{len(stray)} stray formula(s) in a manual column ({examples})", "low"))

    def check_orphans(self) -> None:
        known = set(self.wb.sheetnames)
        for sheet in self.x["sheets"]:
            for unknown in sheet.get("unknown_sheet_refs", []):
                self.findings.append(Finding(
                    "orphan references", f"{sheet['name']} (formulas)",
                    "formula references resolve to sheets in this workbook",
                    f"referenced sheet '{unknown}' does not exist", "high"))
        for nr in self.x.get("workbook", {}).get("named_ranges", []):
            if isinstance(nr, dict) and "#REF!" in str(nr.get("target", "")):
                self.findings.append(Finding(
                    "orphan references", f"named range '{nr.get('name')}'",
                    "named range points at a live range",
                    f"target is broken: {nr.get('target')}", "medium"))
        for (t_sheet, t_col) in sorted(self.lookup_targets):
            if t_sheet not in known:
                continue
            meta = self.sheet_meta(t_sheet)
            if not meta:
                continue
            col = next((c for c in meta["columns"] if c["column"] == t_col), None)
            if col is None or col["rows_nonempty"] == 0:
                self.findings.append(Finding(
                    "orphan references", f"{t_sheet}!{t_col}",
                    "lookup target column holds data",
                    "lookups point at an empty column — every lookup fails", "medium"))


_UNSUPPORTED = object()

CHECK_ORDER = ["lookup integrity", "key uniqueness", "whitespace/casing",
               "staleness", "pattern exceptions", "orphan references"]


def render_replay_md(extraction: dict, rep: Replayer, out_path: str) -> tuple[int, dict]:
    sev_rank = Finding.ORDER
    findings = sorted(rep.findings, key=lambda f: (sev_rank[f.severity], CHECK_ORDER.index(f.check), f.location))
    by_check = {c: Counter() for c in CHECK_ORDER}
    for f in findings:
        by_check[f.check][f.severity] += 1

    lines = []
    ident = extraction["identity"]
    lines.append(f"# Replay report — {ident['filename']}")
    lines.append("")
    lines.append(f"> Source sha256: `{ident['sha256'][:12]}` | Replayed: {now_iso()} | "
                 f"Tool: xlsx_extract.py v{TOOL_VERSION} | Cells checked: {rep.checked_cells} "
                 f"(skipped {rep.skipped_empty} empty lookups)")
    lines.append("> Rule: **no check = no claim.** Every row below is a mechanical recomputation; "
                 "anything the battery cannot replay is listed under *Not replayable*.")
    lines.append("")
    lines.append("## Summary")
    lines.append("")
    lines.append("| check | findings | high | medium | low | info |")
    lines.append("|---|---|---|---|---|---|")
    for check in CHECK_ORDER:
        c = by_check[check]
        total = sum(c.values())
        lines.append(f"| {check} | {total} | {c['high']} | {c['medium']} | {c['low']} | {c['info']} |")
    lines.append("")

    lines.append("## Findings")
    lines.append("")
    if findings:
        lines.append("| # | check | location | expected | found | severity |")
        lines.append("|---|---|---|---|---|---|")
        for i, f in enumerate(findings, 1):
            def esc(s: str) -> str:
                return str(s).replace("|", "\\|").replace("\n", " ")
            lines.append(f"| {i} | {f.check} | `{esc(f.location)}` | {esc(f.expected)} | "
                         f"{esc(f.found)} | **{f.severity}** |")
    else:
        lines.append("**0 findings.** The battery ran and found nothing — absence of findings is "
                     "itself evidence (checked cells: " + str(rep.checked_cells) + ").")
    lines.append("")

    lines.append("## Not replayable")
    lines.append("")
    if rep.not_replayable:
        lines.append("| formula / reason | cells | example |")
        lines.append("|---|---|---|")
        for pattern, count in rep.not_replayable.most_common():
            safe = pattern.replace("|", "\\|")
            lines.append(f"| {safe} | {count} | `{rep.nr_example[pattern]}` |")
        lines.append("")
        lines.append("These are Unknowns for the process model — never inferred.")
    else:
        lines.append("None — every formula family found was inside the supported battery.")
    lines.append("")
    lines.append("## Notes")
    lines.append("")
    lines.append(f"- Staleness ages are computed against the replay date ({dt.date.today().isoformat()}); "
                 "re-running on another day shifts day counts, not conclusions.")
    lines.append("- Text matching mirrors Excel: case-insensitive; the TRIM pass isolates "
                 "whitespace-only defects.")
    lines.append("")

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    tmp = out_path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    os.replace(tmp, out_path)
    return len(findings), {c: sum(by_check[c].values()) for c in CHECK_ORDER}


def replay(path: str, extraction_path: str, out_path: str, log_path: str | None) -> int:
    filename = os.path.basename(path)
    for p in (path, extraction_path):
        if not os.path.exists(p):
            print(f"not found: {p}", file=sys.stderr)
            return 2
    with open(extraction_path, encoding="utf-8") as fh:
        extraction = json.load(fh)
    if extraction.get("status") != "ok":
        print(f"extraction status is '{extraction.get('status')}' — nothing to replay "
              f"(reason: {extraction.get('reason')})", file=sys.stderr)
        return 3
    sha = sha256_file(path)
    if extraction.get("identity", {}).get("sha256") != sha:
        print("extraction is STALE (file sha256 changed) — re-run L1 extraction first", file=sys.stderr)
        append_log(log_path, "L3", filename, "stale-extraction", f"file sha256 {sha[:8]}")
        return 3

    try:
        wb_data = load_workbook(path, data_only=True)
        wb_formula = load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"replay failed to open workbook: {exc}", file=sys.stderr)
        append_log(log_path, "L3", filename, "failed", short(str(exc), 120))
        return 3

    rep = Replayer(wb_data, extraction)
    rep._fwb = wb_formula
    rep.run()
    n, per_check = render_replay_md(extraction, rep, out_path)
    print(f"replayed: {filename} -> {out_path} ({n} findings; "
          + ", ".join(f"{k}: {v}" for k, v in per_check.items() if v) + (")" if any(per_check.values()) else "0 across all checks)"))
    append_log(log_path, "L3", filename, "replayed", f"{n} findings")
    return 0


# --------------------------------------------------------------------- main

def main(argv=None) -> int:
    utf8_console()
    parser = argparse.ArgumentParser(description="aisa process-capture: L1 xlsx extraction + L3 replay")
    parser.add_argument("--replay", action="store_true", help="run the L3 replay battery")
    parser.add_argument("--force", action="store_true", help="ignore extraction cache")
    parser.add_argument("--log", default=None, help="append events to this capture log file")
    parser.add_argument("--version", action="version", version=f"xlsx_extract.py {TOOL_VERSION}")
    parser.add_argument("paths", nargs="+", help="extract: <input.xlsx> <out.json> | "
                                                 "replay: <input.xlsx> <extraction.json> <out.md>")
    args = parser.parse_args(argv)

    if args.replay:
        if len(args.paths) != 3:
            parser.error("--replay needs: <input.xlsx> <extraction.json> <out.replay.md>")
        return replay(args.paths[0], args.paths[1], args.paths[2], args.log)
    if len(args.paths) != 2:
        parser.error("extraction needs: <input.xlsx> <out.extraction.json>")
    return extract(args.paths[0], args.paths[1], args.force, args.log)


if __name__ == "__main__":
    sys.exit(main())
